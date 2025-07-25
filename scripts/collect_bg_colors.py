import os
from openpyxl import load_workbook
from constants import COLOR_MEANINGS, KNOWN_COLORS

VAULT_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
EXCEL_PATH = os.path.join(VAULT_PATH, "resources", "乔剪国战表格.xlsx")

COL_PACKAGE = 1
COL_CODE = 2
COL_NAME = 3
COL_TITLE = 4
COL_HP = 5
COL_SYNERGY = 6
FIRST_6_COLS = [COL_PACKAGE, COL_CODE, COL_NAME, COL_TITLE, COL_HP, COL_SYNERGY]

SKIP_PREFIXES = ("黑桃", "梅花", "红桃", "方片", "EMA", "Z.")
TARGET_SHEETS = [
    "官方范围",
    "兵情篇",
    "豪门贵胄",
    "异军突起",
    "小型扩展",
    "他山之玉"
]

def get_cell_bg_color(cell):
    if cell.fill and cell.fill.fgColor and cell.fill.fgColor.type == "rgb":
        return cell.fill.fgColor.rgb
    return None

def is_valid_wujiang_row(ws, row_idx):
    code_cell = ws.cell(row=row_idx, column=COL_CODE)
    name_cell = ws.cell(row=row_idx, column=COL_NAME)
    if not code_cell.value or not name_cell.value:
        return False
    code = str(code_cell.value).strip()
    if code.startswith(SKIP_PREFIXES):
        return False
    return True

def collect_colors_from_sheet(ws, color_counter, unknown_color_rows, total_wujiang):
    for row_idx in range(2, ws.max_row + 1):
        # 跳过隐藏行
        if ws.row_dimensions[row_idx].hidden:
            continue
        if not is_valid_wujiang_row(ws, row_idx):
            continue

        total_wujiang[0] += 1  # 计数总武将数

        code = str(ws.cell(row=row_idx, column=COL_CODE).value).strip()
        hero_name = str(ws.cell(row=row_idx, column=COL_NAME).value).strip()
        code_name = f"{code} {hero_name}"

        # 前 6 列找颜色
        for col_idx in FIRST_6_COLS:
            color = get_cell_bg_color(ws.cell(row=row_idx, column=col_idx))
            if not color or color in ("00000000", "FFFFFFFF"):
                continue

            if color in KNOWN_COLORS:
                # ✅ 已知 9 种颜色 → 计数
                color_counter[color] = color_counter.get(color, 0) + 1
            else:
                # ✅ 未知颜色 → 保存具体行
                if color not in unknown_color_rows:
                    unknown_color_rows[color] = []
                unknown_color_rows[color].append(code_name)

def main():
    wb = load_workbook(EXCEL_PATH)
    total_wujiang = [0]  # 用 list 方便引用计数
    color_counter = {}   # 已知颜色 → 次数
    unknown_color_rows = {}  # 未知颜色 → 行列表

    for sheet_name in TARGET_SHEETS:
        ws = wb[sheet_name]
        collect_colors_from_sheet(ws, color_counter, unknown_color_rows, total_wujiang)

    # 输出总数
    print(f"\n📊 武将总数: {total_wujiang[0]}")

    # 输出已知 9 色统计
    print("\n🎨 已知颜色统计：")
    for color, meaning in COLOR_MEANINGS.items():
        count = color_counter.get(color, 0)
        print(f"- {meaning} ({color}) : {count} 次")

    # 输出未知颜色及所有行
    if unknown_color_rows:
        print("\n❓ 未知颜色及行：")
        for color, rows in unknown_color_rows.items():
            print(f"\n未知颜色 {color} ({len(rows)} 行)")
            for r in rows:
                print(f"  - {r}")
    else:
        print("\n✅ 没有发现未知颜色")

if __name__ == "__main__":
    main()
